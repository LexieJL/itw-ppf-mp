# 从 【4国语言翻译】模版-小程序用2.xlsx 生成 utils/catalog.js 和 images/brands/*.png
# 用法：新建一个工作目录 D，把产品表复制成 D/src.xlsx，并解压到 D/xl/（unzip D/src.xlsx -d D/xl）取出 logo，然后运行 python3 tools/gen_catalog.py D
# 依赖 openpyxl、Pillow；LOGO 映射按品牌介绍表的行序人工核对过，表格行序变了要重新核对
import openpyxl, json, re, os, sys
from PIL import Image
S = sys.argv[1] if len(sys.argv) > 1 else os.path.dirname(os.path.abspath(__file__))
MP = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))  # miniprogram/ 目录
wb = openpyxl.load_workbook(os.path.join(S, 'src.xlsx'))

# 品牌介绍表顺序 → 对应 drawing 里的图片（已人工核对）
LOGO = {1:1,2:2,3:3,4:4,5:5,6:6,7:7,8:8,9:9,11:10,12:11,13:12,14:13,15:14,16:15,17:20,18:23,19:21,20:22,21:16,22:17,23:18,24:19}
def key(s): return re.sub(r"[^a-z]", '', (s or '').lower())
ALIAS = {'wynns': 'wynns', 'epocast36': 'epocast', 'accululbe': 'acculube'}
brands = {}
for i, r in enumerate(list(wb['品牌介绍'].iter_rows(values_only=True))[1:], 1):
    if not r[0]: continue
    k = key(r[0]); cn = (r[1] or '').strip()
    b = {'id': k, 'name': r[0].strip(), 'cn': '' if cn in ('', '--') else cn,
         'intro': (r[2] or '').strip(), 'logo': '', 'order': i}
    if i in LOGO:
        im = Image.open(os.path.join(S, 'xl/xl/media/image%d.png' % LOGO[i])).convert('RGBA')
        im.thumbnail((240, 96))
        os.makedirs(MP + '/images/brands', exist_ok=True)
        im.save(MP + '/images/brands/%s.png' % k, optimize=True)
        b['logo'] = '/images/brands/%s.png' % k
    brands[k] = b
# CFO / CFR 是 Chockfast 系列（中文名 巧固快），Chockfast 行本身没有 logo，借用 CFO 的
brands['chockfast']['logo'] = brands['cfo']['logo']
if not brands['chockfast']['cn']: brands['chockfast']['cn'] = brands['cfo']['cn']

rows = list(wb['CN'].iter_rows(values_only=True))[1:]
prods, order = {}, []
last = None
def clean(s): return (str(s).strip() if s is not None else '')
for r in rows:
    if not any(r): continue
    brand = clean(r[1]) or (last and last[1])  # 一行品牌留空，沿用上一行
    if not clean(r[1]) and last: r = tuple(x if x is not None else (last[j] if j in (1,3) else None) for j, x in enumerate(r))
    last = r
    k = key(brand); k = ALIAS.get(k, k)
    if k not in brands:
        brands[k] = {'id': k, 'name': brand, 'cn': '', 'intro': '', 'logo': '', 'order': 99}
    name = re.sub(r'\s+', ' ', clean(r[5]))
    pk = (k, name.lower())
    if pk not in prods:
        prods[pk] = {'brand': k, 'series': clean(r[4]) or '其他', 'name': name,
                     'markets': [], 'fields': [], 'intro': '', 'usage': '', 'skus': [],
                     'inBook': '', 'isNew': False}
        order.append(pk)
    p = prods[pk]
    m = clean(r[2])
    if m and m not in p['markets']: p['markets'].append(m)
    for f in clean(r[3]).replace('，', ',').split(','):
        f = f.strip()
        if f and f not in p['fields']: p['fields'].append(f)
    if not p['intro']: p['intro'] = clean(r[10]) or clean(r[9])
    if not p['usage']: p['usage'] = clean(r[11])
    code, size = clean(r[6]), clean(r[8])
    if (code or size) and not any(s['code'] == code and s['size'] == size for s in p['skus']):
        p['skus'].append({'code': code, 'size': size, 'img': clean(r[7])})
    if clean(r[12]) == '是': p['inBook'] = '是'
    elif clean(r[12]).startswith('可删除') and p['inBook'] != '是': p['inBook'] = '可删除'
    if clean(r[13]): p['isNew'] = True

counts = {}
for pk in order: counts[prods[pk]['brand']] = counts.get(prods[pk]['brand'], 0) + 1
blist = sorted(brands.values(), key=lambda b: (-counts.get(b['id'], 0), b['order']))
out_b = [{k: b[k] for k in ('id', 'name', 'cn', 'logo', 'intro')} for b in blist]
out_p = []
for n, pk in enumerate(order, 1):
    p = prods[pk]; p = dict(id='p%d' % n, **p)
    out_p.append(p)
src = '// 由 ITW 提供的 【4国语言翻译】模版-小程序用2.xlsx 自动生成（目前只有中文），请勿手改；正式上线后改从 CloudBase 读取\n'
src += '// BRANDS：品牌介绍表 + CN 表里出现的品牌；PRODUCTS：CN 表按 品牌+产品名 合并，同一产品的不同代码/规格放在 skus\n'
src += 'const BRANDS = ' + json.dumps(out_b, ensure_ascii=False, indent=0) + '\n\n'
src += 'const PRODUCTS = ' + json.dumps(out_p, ensure_ascii=False, separators=(',', ':')).replace('},{"id"', '},\n{"id"') + '\n\n'
src += 'module.exports = { BRANDS, PRODUCTS }\n'
open(MP + '/utils/catalog.js', 'w').write(src)
print(len(out_b), 'brands', len(out_p), 'products', os.path.getsize(MP + '/utils/catalog.js'), 'bytes')
for b in out_b: print(b['name'], b['cn'], counts.get(b['id'], 0), 'logo' if b['logo'] else '-', len(b['intro']))
from collections import Counter
print(Counter(p['inBook'] for p in out_p), sum(p['isNew'] for p in out_p))
print(sum(len(p['skus']) for p in out_p), 'skus')
